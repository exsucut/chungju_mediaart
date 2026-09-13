# -*- coding: utf-8 -*-
"""확정 컷으로 PPTX 를 만든다 — 한 컷 한 장.

보드의 PDF 내보내기는 브라우저 인쇄라 손을 못 댄다. 이건 파워포인트 파일이라
클라이언트가 슬라이드 위에서 바로 메모하고 잘라 쓸 수 있다.

    python build_web.py && python make_pptx.py
    python make_pptx.py --only S5 S6 S7     # 일부만

한 장에 들어가는 것
    상단   컷 ID · 이름 · 타임코드 · 길이
    큰 그림 캔버스 전체(5320x1600) + 좌/우 화면 영역 · 갭 · 꺾임선 표시
    아래   좌측 스크린 크롭 · 우측 파사드 크롭 · 바닥 1:1
    본문   레이아웃 · 카메라 · 설명 · 전환   — 전부 편집 가능한 텍스트

그림은 PIL 로 미리 합성해 넣는다. 도형으로 그리면 파워포인트에서 위치가
틀어지기 쉽다.
"""
import argparse
import io
import json
import os
import re
import sys
import tempfile

from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "images")

# 화면 배치 — build_web.py 와 같은 값을 쓴다
GAP, OFFSET_L = 200, 640
SCR = {"L": dict(w=1920, h=960, x=0, y=OFFSET_L),
       "R": dict(w=3200, h=1200, x=1920 + GAP, y=0)}
CANVAS_W = SCR["R"]["x"] + SCR["R"]["w"]                    # 5320
CANVAS_H = max(SCR["L"]["y"] + SCR["L"]["h"], SCR["R"]["h"])  # 1600
FOLD_X = int(SCR["R"]["x"] + SCR["R"]["w"] * 0.58)
CLEAR_L, CLEAR_R = 1850, 2210
ALIGN_Y = 0.62

RED, BLUE, YELLOW, INK = (226, 74, 74), (74, 142, 226), (232, 196, 42), (24, 26, 30)


def _find_out():
    p = HERE
    for _ in range(4):
        p = os.path.dirname(p)
        if os.path.exists(os.path.join(p, "index.html")):
            return p
    return os.path.dirname(os.path.dirname(HERE))


OUT = _find_out()


def pad_sid(sid):
    m = re.match(r"^S(\d+)([a-z]*)$", sid)
    return f"S{int(m.group(1)):02d}{m.group(2)}" if m else sid


def canvas(path, align):
    """플레이트를 캔버스 폭에 맞춰 얹고 남는 세로를 align 비율로 버린다."""
    im = Image.open(path).convert("RGB")
    ph = int(round(im.height * CANVAS_W / im.width))
    im = im.resize((CANVAS_W, ph), Image.LANCZOS)
    c = Image.new("RGB", (CANVAS_W, CANVAS_H), (0, 0, 0))
    c.paste(im, (0, -int(round((ph - CANVAS_H) * align))))
    return c


def marked(c):
    """좌/우 화면 · 갭 · 꺾임선 · 여백 존을 그려 넣은 사본."""
    m = c.copy()
    d = ImageDraw.Draw(m)
    r = SCR["L"]; d.rectangle([r["x"], r["y"], r["x"] + r["w"] - 1, r["y"] + r["h"] - 1],
                              outline=RED, width=7)
    r = SCR["R"]; d.rectangle([r["x"], r["y"], r["x"] + r["w"] - 1, r["y"] + r["h"] - 1],
                              outline=BLUE, width=7)
    d.rectangle([CLEAR_L, 0, CLEAR_R, CANVAS_H - 1], outline=YELLOW, width=6)
    d.line([FOLD_X, 0, FOLD_X, CANVAS_H - 1], fill=YELLOW, width=5)
    return m


def crop(c, tag):
    r = SCR[tag]
    return c.crop((r["x"], r["y"], r["x"] + r["w"], r["y"] + r["h"]))


def save_tmp(im, tmpdir, name, w=None):
    if w and im.width > w:
        im = im.resize((w, max(1, int(im.height * w / im.width))), Image.LANCZOS)
    p = os.path.join(tmpdir, name)
    im.convert("RGB").save(p, "JPEG", quality=88, optimize=True)
    return p


def build(only=None, dest=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    shots = json.load(io.open(os.path.join(HERE, "shots.json"), encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    blank = prs.slide_layouts[6]
    SW, SH = prs.slide_width, prs.slide_height
    M = Emu(380000)                                    # 여백

    def textbox(sl, x, y, w, h, paras):
        """paras = [[(글자, 크기, 굵게, 흐리게), ...], ...] — 안쪽 리스트가 한 줄."""
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            for txt, size, bold, grey in runs:
                r = p.add_run(); r.text = txt
                r.font.size = Pt(size); r.font.bold = bold
                r.font.color.rgb = (RGBColor(0x7A, 0x83, 0x8C) if grey
                                    else RGBColor(0x18, 0x1A, 0x1E))
                r.font.name = "맑은 고딕"
        return tb

    tmpdir = tempfile.mkdtemp(prefix="sbpptx_")
    n = 0
    for act in shots["acts"]:
        for s in act["shots"]:
            if only and s["id"] not in only:
                continue
            vs = (s.get("variants") or s.get("variantsR") or s.get("variantsL") or [])
            if not vs:
                continue
            f = vs[0]["f"]
            p = os.path.join(SRC, f.replace("/", os.sep))
            if not os.path.exists(p):
                print(f"  ! 원본 없음: {f}", file=sys.stderr)
                continue
            align = vs[0].get("align", s.get("align", ALIGN_Y))
            c = canvas(p, align)

            sl = prs.slides.add_slide(blank)

            FULL = SW - 2 * M
            RW = Emu(3400000)                      # 아래 오른쪽 글 칸
            COLGAP = Emu(260000)

            # ── 머리 ─────────────────────────────────────────────
            HEAD_Y, HEAD_H = Emu(190000), Emu(330000)
            textbox(sl, M, HEAD_Y, FULL, HEAD_H, [[
                (f"{s['id']}   ", 20, True, False),
                (f"{s.get('name','')}    ", 15, True, False),
                (f"{s.get('tc','')} · {s.get('len','')}", 10, False, True)]])

            # ── 큰 그림 — 전체 폭 ────────────────────────────────
            top = HEAD_Y + HEAD_H + Emu(90000)
            bh = int(FULL * CANVAS_H / CANVAS_W)
            sl.shapes.add_picture(save_tmp(marked(c), tmpdir, f"{n}_big.jpg", 2600),
                                  M, top, width=FULL, height=bh)
            cap_y = top + bh + Emu(55000)
            textbox(sl, M, cap_y, FULL, Emu(200000), [[
                (f"좌 1920×960 · 우 3200×1200 · 갭 200px = 실물 철당간 · 꺾임선 74.7% "
                 f"· 세로 위치 {align:.2f}", 8, False, True)]])

            # ── 아래 — 왼쪽 크롭 석 장 / 오른쪽 글 ────────────────
            cells = [("좌측 스크린 · 2:1", crop(c, "L")),
                     ("우측 파사드 · 8:3", crop(c, "R"))]
            fl = (s.get("floor") or [])
            if fl:
                fp = os.path.join(SRC, fl[0]["f"].replace("/", os.sep))
                if os.path.exists(fp):
                    cells.append((f"바닥 1:1 · {fl[0].get('label','')}",
                                  Image.open(fp).convert("RGB")))
            gap, cap_h = Emu(120000), Emu(190000)
            row_y = cap_y + Emu(280000)
            thumbs_w = FULL - RW - COLGAP
            ratio = sum(im.width / im.height for _, im in cells)
            th = int((thumbs_w - gap * (len(cells) - 1)) / ratio)
            room = SH - M - cap_h - row_y          # 아래로 넘치지 않게
            if th > room:
                th = room
            x = M
            for i, (lab, im) in enumerate(cells):
                iw = int(th * im.width / im.height)
                sl.shapes.add_picture(save_tmp(im, tmpdir, f"{n}_{i}.jpg", 1500),
                                      x, row_y, width=iw, height=th)
                textbox(sl, x, row_y + th + Emu(30000), iw, cap_h,
                        [[(lab, 8, False, True)]])
                x += iw + gap

            # ── 글 칸 ────────────────────────────────────────────
            bx = M + thumbs_w + COLGAP
            paras = []
            for k, lab in (("layout", "레이아웃"), ("cam", "카메라"),
                           ("desc", "설명"), ("trans", "전환")):
                if s.get(k):
                    paras.append([(lab, 8, True, True)])
                    paras.append([(s[k], 9, False, False)])
                    paras.append([("", 4, False, False)])
            if paras:
                textbox(sl, bx, row_y, RW, SH - M - row_y, paras)
            n += 1
            print(f"  {s['id']:5s} {s.get('name','')}")

    if not n:
        print("담을 컷이 없다.", file=sys.stderr)
        return 1
    dest = dest or os.path.join(OUT, "제오경_스토리보드_확정.pptx")
    prs.save(dest)
    print(f"\nPPTX {n}장  {os.path.getsize(dest)/1024/1024:.1f} MB")
    print(f"  → {dest}")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", nargs="*", help="컷 ID 만 골라 담는다")
    ap.add_argument("--out", help="저장 경로")
    a = ap.parse_args()
    sys.exit(build(set(a.only) if a.only else None, a.out))
