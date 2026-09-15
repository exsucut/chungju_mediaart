# -*- coding: utf-8 -*-
"""2차 피드백 회신을 PPTX 로 만든다.

클라이언트가 PPT 로 피드백을 줬으니 회신도 같은 형식으로 간다. 마크다운
문서(`client/답변_2차피드백_고증및질문_2026-09-15.md`)가 원문이고, 이 스크립트는
그걸 슬라이드로 옮기면서 **전/후 그림을 붙인다.** 글만 있는 회신보다
"뭘 어떻게 고쳤는지"가 한 장에서 읽힌다.

    python make_reply_pptx.py

구성
    표지 · 컷번호 대조표
    이번에 반영한 것        — 전/후 그림 한 장씩 (5장)
    고증 근거 자료          — 무엇을 보고 그렸나
    고증 점검 · 판단 필요    — 읍성 형식 · 시장 · 물속 유물 편년
    질문 4건 답변
    지시를 기다리는 것 · 참고 출처

그림은 make_pptx.py 와 같은 캔버스 규격(5320x1600)으로 합성한다.
"""
import io
import os
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "images")
OUT = os.path.dirname(os.path.dirname(HERE))

CANVAS_W, CANVAS_H = 5320, 1600
ALIGN_Y = 0.62

INK = RGBColor(0x18, 0x1A, 0x1E)
GREY = RGBColor(0x76, 0x7E, 0x88)
ACCENT = RGBColor(0xB4, 0x5A, 0x28)      # 고친 것 · 강조
WARN = RGBColor(0x8A, 0x6D, 0x1E)        # 판단 필요
RULE = RGBColor(0xD8, 0xDC, 0xE2)

SW, SH = Emu(12192000), Emu(6858000)
M = Emu(560000)
FULL = SW - 2 * M


def canvas(name, align=ALIGN_Y):
    """플레이트를 캔버스 폭에 맞춰 얹고 남는 세로를 align 비율로 버린다."""
    im = Image.open(os.path.join(SRC, name)).convert("RGB")
    ph = int(round(im.height * CANVAS_W / im.width))
    im = im.resize((CANVAS_W, ph), Image.LANCZOS)
    c = Image.new("RGB", (CANVAS_W, CANVAS_H), (0, 0, 0))
    c.paste(im, (0, -int(round((ph - CANVAS_H) * align))))
    return c


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]
        self.tmp = tempfile.mkdtemp(prefix="reply_")
        self.n = 0

    # ── 기본 조각 ────────────────────────────────────────────
    def slide(self):
        return self.prs.slides.add_slide(self.blank)

    def text(self, sl, x, y, w, h, paras, space=6):
        """paras = [[(글자, pt, 굵게, 색), ...], ...] — 안쪽 리스트가 한 줄."""
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space)
            for txt, size, bold, col in runs:
                r = p.add_run()
                r.text = txt
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.color.rgb = col
                r.font.name = "맑은 고딕"
        return tb

    def head(self, sl, kicker, title):
        y = Emu(420000)
        if kicker:
            self.text(sl, M, y, FULL, Emu(230000),
                      [[(kicker, 10, True, GREY)]])
            y += Emu(280000)
        self.text(sl, M, y, FULL, Emu(400000), [[(title, 22, True, INK)]])
        ln = sl.shapes.add_shape(1, M, y + Emu(500000), FULL, Emu(9000))
        ln.fill.solid()
        ln.fill.fore_color.rgb = RULE
        ln.line.fill.background()
        ln.shadow.inherit = False
        return y + Emu(700000)

    def foot(self, sl, txt):
        self.text(sl, M, SH - Emu(500000), FULL, Emu(220000),
                  [[(txt, 8, False, GREY)]])

    def pic(self, sl, im, x, y, w, cap=None, px=2400):
        if im.width > px:
            im = im.resize((px, int(im.height * px / im.width)), Image.LANCZOS)
        p = os.path.join(self.tmp, f"{self.n}.jpg")
        self.n += 1
        im.convert("RGB").save(p, "JPEG", quality=87, optimize=True)
        h = int(w * im.height / im.width)
        sl.shapes.add_picture(p, x, y, width=w, height=h)
        if cap:
            self.text(sl, x, y + h + Emu(70000), w, Emu(200000),
                      [[(cap, 9, True, GREY)]])
        return h

    def table(self, sl, x, y, w, rows, widths, size=10, head=True):
        """rows = [[셀, 셀, ...], ...]. widths 는 비율."""
        tot = sum(widths)
        cy = y
        for ri, row in enumerate(rows):
            cx = x
            tall = Emu(0)
            for ci, cell in enumerate(row):
                cw = int(w * widths[ci] / tot)
                bold = head and ri == 0
                col = GREY if bold else INK
                tb = self.text(sl, cx + Emu(60000), cy + Emu(50000),
                               cw - Emu(120000), Emu(260000),
                               [[(cell, size, bold, col)]], space=0)
                est = Emu(int(260000 + 150000 * (len(cell) // max(1, int(cw / Emu(95000))))))
                tall = max(tall, est)
                cx += cw
            cy += tall + Emu(60000)
            if head and ri == 0:
                ln = sl.shapes.add_shape(1, x, cy - Emu(40000), w, Emu(6000))
                ln.fill.solid()
                ln.fill.fore_color.rgb = RULE
                ln.line.fill.background()
                ln.shadow.inherit = False
        return cy


def build():
    d = Deck()

    # ══ 표지 ════════════════════════════════════════════════
    sl = d.slide()
    d.text(sl, M, Emu(2100000), FULL, Emu(400000),
           [[("제오경 「염원 – 주성의 돛대」", 13, True, GREY)]])
    d.text(sl, M, Emu(2500000), FULL, Emu(700000),
           [[("2차 피드백 회신", 36, True, INK)]])
    d.text(sl, M, Emu(3350000), FULL, Emu(900000), [
        [("반영 내역 · 고증 근거 · 질문 4건 답변", 15, False, INK)],
        [("", 6, False, INK)],
        [("2026. 09. 16.    기준 문서 260914_청주국가유산미디어아트_피드백.pptx", 10, False, GREY)],
    ])

    # ══ 컷 번호 대조 ═════════════════════════════════════════
    sl = d.slide()
    y = d.head(sl, "0", "먼저 — 컷 번호가 서로 다릅니다")
    d.text(sl, M, y, FULL, Emu(400000), [
        [("보내주신 PPT 와 저희 스토리보드의 번호가 다릅니다. 섞어 쓰면 엉뚱한 컷을 "
          "고치게 되어, 앞으로 이 대조표를 기준으로 쓰겠습니다. 아래 본문은 ", 11, False, INK),
         ("PPT 번호(저희 번호)", 11, True, INK), (" 로 병기합니다.", 11, False, INK)]])
    y += Emu(560000)
    pairs = [("S02", "S0", "안개 속 배의 선미"), ("S03", "S1", "안개가 먹구름으로"),
             ("S04", "S1b", "먹구름 걷히고 멀리 땅"), ("S05 A/B", "S2a / S2b", "도성 부감"),
             ("S06", "S3", "도성 안으로 급하강"), ("S07 · S08", "S4a", "장터"),
             ("S09", "S4b", "냇가의 아이들"), ("S10", "S4c", "지게 진 일꾼과 언덕"),
             ("S11", "S4d", "절 문"), ("S12", "S5", "용두사 전경"),
             ("S13", "S6", "하늘에 먹구름"), ("S14", "S7", "폭우"),
             ("S15", "S8a", "파도가 들이침"), ("S16", "S8b", "큰 파도가 덮침"),
             ("S17", "S9", "침수 · 수면 아래"), ("S18", "S11 · S10 · S12", "물속 유물")]
    half = (len(pairs) + 1) // 2
    cw = int(FULL / 2) - Emu(200000)
    for col, chunk in enumerate((pairs[:half], pairs[half:])):
        rows = [["PPT", "저희", "내용"]] + [list(p) for p in chunk]
        d.table(sl, M + col * (cw + Emu(400000)), y, cw, rows, [18, 26, 56], size=10)
    d.foot(sl, "PPT S18 한 장이 저희 쪽에서는 세 컷(S11 빛 → S10 부유 → S12 산란)으로 나뉩니다.")

    # ══ 반영한 것 — 전/후 ════════════════════════════════════
    REWORK = [
        dict(ppt="PPT S02  (저희 S0)", title="배가 서양 배로 읽힌다",
             quote="“배의 상부 파이프가 너무 잘 보여서 서양 배인게 확연하게 드러남. 좀더 아래에서 "
                   "위로 선미를 찍어서 윗부분이 잘 안 드러나게 했으면 하면서 동시에, 입체감이 더 "
                   "도드라졌으면 함. 위에 안개를 더 씌우는 것도 좋음”",
             before="S00_안개의배_클라이언트확정_r8.jpg", after="S00_안개의배_로우앵글_r44_va.png",
             done=[("카메라를 수면까지 내려 ", False), ("선미를 올려다보게", True),
                   (" 했습니다. 상부 구조는 안개로 덮고 파이프·펀넬·철제 난간을 전부 뺐습니다. "
                    "측광으로 선체 볼륨을 세워 ", False), ("입체감", True), ("도 같이 잡았습니다.", False)]),
        dict(ppt="PPT S12  (저희 S5)", title="용두사 전경 — 바닥이 휑하다",
             quote="“바닥이 너무 휑해 보여서 잔디 조금만 깔아주세요” · "
                   "“라이트를 조금 더 밝게 하여 먹구름과 대비를 강조해주세요”",
             before="S05_용두사_포토샵확정_r26.png", after="S05_용두사_잔디_맑음_r45_va.png",
             done=[("마당에 ", False), ("마모된 잔디", True),
                   (" 를 깔고 볕을 세웠습니다. 밝기를 전체로 올리면 앞뒤 컷과 벌어지므로 "
                    "올린 것은 ", False), ("땅에 드는 볕뿐", True),
                   ("이고, 하늘은 확정본의 맑은 하늘 그대로 뒀습니다. "
                    "먹구름과의 대비는 다음 컷(S6)에서 만들어집니다.", False)]),
        dict(ppt="PPT S13  (저희 S6)", title="하늘에 먹구름 — 좌에서 우로",
             quote="“구름의 무빙이 왼쪽에서 오른쪽으로 흘러지고, 바닥도 왼쪽에서 오른쪽으로 "
                   "어두워지게 타이밍을 맞춰주세요”",
             before="S06_비의시작_S5확정_r27_va.png", after="S06_비의시작_좌우그늘_r45_va.png",
             done=[("같은 잔디와 볕을 유지한 채 구름을 좌측에서 밀어넣었습니다. "
                    "지금은 ", False), ("좌측이 이미 그늘, 우측은 아직 볕", True),
                   (" 인 중간 상태입니다 — 좌 25% / 우 51% 로 방향이 스틸에서도 읽힙니다. "
                    "실제 무빙 타이밍은 영상 단계에서 이 방향 그대로 잡겠습니다.", False)]),
        dict(ppt="PPT S14  (저희 S7)", title="폭우 — 환경이 또렷하면 안 된다",
             quote="“폭우가 칠때는 또렷하게 환경이 보이면 안될 것 같습니다. 프레임(화면)위 물이 "
                   "뭍어난 느낌을 주거나, 환경자체를 뿌옇게 하는 방안” · "
                   "“번개가 더 잘 보이게 환경이 조금 더 어두울 수 있을까요?”",
             before="S07_폭우_S5확정_r27_va.png", after="S07_폭우_렌즈물_r50.png",
             done=[("두 가지를 다 했습니다. ", False), ("① 환경을 지웠습니다", True),
                   (" — 비의 장막과 물보라를 카메라와 대상 사이에 한 겹 더 세워, 먼 산은 완전히 "
                    "사라지고 건물은 지붕이라는 것만 알아볼 수 있는 덩어리로 남습니다. "
                    "서까래·공포·난간살·기왓골은 화면 어디에도 읽히지 않습니다. ", False),
                   ("② 프레임 위의 물", True),
                   (" — 렌즈 앞에 맺힌 물방울과 흘러내린 자국을 가장자리에 얹고 가운데는 얇게 "
                    "두어 그림은 읽히게 했습니다. 밝기는 22% 로 한 단계 더 내렸고 번개는 "
                    "그대로 세게 읽힙니다.", False)]),
        dict(ppt="자체 고증 점검  (저희 S2b · S3)", title="지붕 비율 — 기와가 너무 많았다",
             quote="『고려도경』 민거조 — “집이 벌집이나 개미집 같고 띠풀로 지붕을 이었다. "
                   "열 집에 한두 집 정도만 기와를 덮었다”",
             before="S02b_주성근접_S2a연결_r29_va.png", after="S02b_주성근접_S2a정합_r49.png",
             done=[("지시받은 항목은 아니지만 실장님이 시장 고증을 짚어주신 김에 전체를 다시 "
                    "훑었습니다. 저희 읍성 컷은 기와가 3분의 1 가까이 차지하고 있었습니다. ", False),
                   ("초가 8 : 기와 2", True),
                   (" 로 다시 잡고 기와는 관아 구획과 문루 등 큰 건물에만 남겼습니다. "
                    "S2a 는 확정해주신 판이라 손대지 않았고, 톤은 계산으로 맞춰 "
                    "S2a → S2b → S3 가 한 호흡으로 이어집니다.", False)]),
    ]

    for i, r in enumerate(REWORK):
        sl = d.slide()
        y = d.head(sl, f"1 — 이번에 반영한 것  ({i + 1}/{len(REWORK)})   {r['ppt']}", r["title"])
        d.text(sl, M, y, FULL, Emu(560000), [[(r["quote"], 10, False, GREY)]])
        y += Emu(300000) + Emu(int(150000 * (len(r["quote"]) // 110 + 1)))
        iw = int(FULL / 2) - Emu(160000)
        h1 = d.pic(sl, canvas(r["before"]), M, y, iw, "기존", px=1800)
        d.pic(sl, canvas(r["after"]), M + iw + Emu(320000), y, iw, "수정", px=1800)
        d.text(sl, M, y + h1 + Emu(460000), FULL, Emu(900000),
               [[(t, 11, b, ACCENT if b else INK) for t, b in r["done"]]])

    # ══ 고증 근거 ═══════════════════════════════════════════
    sl = d.slide()
    y = d.head(sl, "2 — 질문에 대한 답  ·  PPT S07", "“고려시대 마을, 어떤 자료를 참고했는지”")
    rows = [["자료", "시기", "무엇을 가져왔나"],
            ["『고려도경(高麗圖經)』 민거조 — 송 사신 서긍, 1123년", "고려",
             "민가의 밀집도와 지붕 비율. “열 집에 한두 집 정도만 기와를 덮었다”"],
            ["『고려도경』 복식조", "고려",
             "평민 복식 — 흰 모시 백저포(白紵袍), 남녀 구분 없이 같은 형태"],
            ["용두사지 철당간 명문(龍頭寺址鐵幢竿記)", "962년",
             "이 작품의 시간 좌표. 준풍 3년 임술년 주성, 김예종 발원 · 김희일 완성. "
             "20단 철통, 명문이 있는 철당간으로는 국내 유일"],
            ["「청주읍성도」 (구례 운조루 소장)", "조선 후기",
             "읍성 안 건물 배치와 민가 분포의 형태 참고 — 시대 문제는 다음 장"],
            ["청주읍성 기록 — 둘레 1,350보, 높이 8자, 여장 566첩, 포루 8, 우물 12", "조선",
             "성문과 성벽의 비례"],
            ["고려청자 편년", "고려",
             "10세기 초 중국 월주요 기술 유입으로 제작 시작. 순화 3년(992)명 고배가 10세기 "
             "유품 — 962년에는 청자가 “막 시작된” 단계라는 것"]]
    y = d.table(sl, M, y, FULL, rows, [34, 10, 56], size=10)
    d.text(sl, M, y + Emu(200000), FULL, Emu(400000),
           [[("실물 자산 — ", 10, True, GREY),
             ("용두사지 철당간 실물 사진(S18b 및 합성용 3D 요소의 형태 기준), "
              "보내주신 유물 렌더 21점(물속 유물 컷의 직접 레퍼런스)", 10, False, INK)]])

    # ══ 고증 점검 ═══════════════════════════════════════════
    sl = d.slide()
    y = d.head(sl, "2 — 고증 점검", "지금 그림에서 걸리는 것")
    d.text(sl, M, y, FULL, Emu(400000),
           [[("실장님이 “시장 물건이 고증에 걸릴 수 있다”고 짚어주신 게 정확했습니다. "
              "그 김에 전체를 다시 훑었더니 같은 성격의 문제가 몇 개 더 있습니다. "
              "판단이 필요한 것과 저희가 이미 고친 것을 나눠 적습니다.", 11, False, INK)]])
    y += Emu(620000)
    items = [
        ("② 판단이 필요합니다", "읍성이 조선 형식입니다", WARN,
         "고려의 지방 읍성은 규모가 작은 토축(土築)이 일반적이었고, 석축에 옹성·치성·해자를 "
         "갖추며 커지는 것은 조선 — 특히 세종 이후입니다. 지금 저희 읍성은 석축 성벽에 홍예 "
         "문루로, 「청주읍성도」를 따랐으니 자연히 조선 형식입니다. 962년으로 보면 어긋납니다.\n"
         "A. 그대로 간다 — 「청주읍성도」는 청주 시민에게 가장 익숙한 청주의 얼굴이고, 미디어아트는 "
         "사료 재현이 아니라 상징을 다룹니다. 「청주」로 바로 읽히는 값이 고증보다 큽니다.\n"
         "B. 토축으로 바꾼다 — 고증은 맞으나 「청주읍성」이라는 인상은 옅어집니다.\n"
         "저희는 A 를 권합니다. 다만 의도한 선택이라는 점을 기록해 두는 것이 나중에 가장 "
         "단단한 방어가 됩니다."),
        ("③ 실장님 지적이 맞습니다", "시장", WARN,
         "고려의 지방 거래는 향시(鄕市)·허시(墟市) 수준이었습니다. 정기 5일장이 자리잡는 건 "
         "조선 후기로, 15세기 후반에 장시가 서기 시작해 18세기에 통일됩니다. 즉 962년에 상설 "
         "시장 거리는 없었습니다. 지금 S4a 는 좌판과 차양이 늘어선 상설 장터로 보입니다.\n"
         "“최대한 빠르게 지나갔으면 좋겠다”는 판단이 가장 현실적인 해법이고 저희도 같은 "
         "의견입니다. 다만 차양과 고정 좌판을 줄이고 땅에 편 멍석과 등짐 보따리 위주로 바꾸면 "
         "“상설 시장”이 아니라 “사람이 모인 날”로 읽혀 위험이 더 줄어듭니다."),
        ("④ 판단이 필요합니다", "물속 유물의 편년", WARN,
         "보내주신 유물 렌더 21점 중 일부는 12–13세기 유물입니다. 상감 문양이 든 청자 매병은 "
         "상감청자라 12세기 이후이고, 나전 소반도 고려 후기에 화려해집니다. 962년에는 청자가 "
         "이제 막 만들어지기 시작한 단계로 기형이 거칠고 종류도 적습니다.\n"
         "다만 이건 관객이 알아보기 어려운 층위이고, 유물의 아름다움이 그 컷의 정서를 만들고 "
         "있습니다. 보내주신 렌더를 그대로 쓸지, 초기 청자·질그릇 위주로 교체할지 지시 "
         "부탁드립니다."),
        ("⑤ 반영하겠습니다", "평민 의복", ACCENT,
         "『고려도경』 기준으로 흰 모시 백저포, 남녀 구분 없이 같은 형태입니다. 현재 시장 컷의 "
         "인물은 대부분 뒷모습이나 역광 실루엣이고 색도 흰색·미색 계열로 맞춰져 있습니다.\n"
         "다만 우측 전경의 옹기장이 한 명은 얼굴과 머리쓰개가 또렷이 읽히고, 그 머리쓰개가 "
         "조선식 망건·탕건에 가깝습니다. 고려 평민은 건(巾)을 두르거나 맨상투에 가까우므로 이 "
         "인물은 다음 판에서 고치겠습니다."),
    ]
    cw = int(FULL / 2) - Emu(240000)
    for i, (tag, title, col, body) in enumerate(items):
        x = M + (i % 2) * (cw + Emu(480000))
        yy = y + (i // 2) * Emu(2350000)
        paras = [[(tag, 10, True, col)], [(title, 14, True, INK)], [("", 4, False, INK)]]
        paras += [[(ln, 10, False, INK)] for ln in body.split("\n")]
        d.text(sl, x, yy, cw, Emu(2200000), paras)
    d.foot(sl, "① 지붕 비율은 앞 장에서 이미 고쳤습니다.")

    # ══ 질문 4건 ════════════════════════════════════════════
    QA = [
        ("Q1", "“시장, 냇가, 지게 씬의 트랜지션이 궁금합니다”  (PPT S09)",
         [("세 컷을 각각 다른 방식으로 잇는 것을 제안드립니다. 같은 방식이 반복되면 지루해지고, "
           "이 구간은 마을을 훑는 흐름이라 리듬이 필요합니다.", False)],
         [["구간", "방식", "내용"],
          ["장터 → 냇가  (S07·S08 → S09)", "물 매치컷",
           "장터에서 물동이의 물이 기울어 쏟아지는 순간, 그 물줄기가 그대로 개천의 흐름이 된다. "
           "형태와 방향을 맞춰 붙이면 컷이 끊기지 않는다"],
          ["냇가 → 지게  (S09 → S10)", "시선 유도 · 오브젝트 팔로우",
           "아이가 띄운 나뭇잎 배가 물길을 타고 흘러가고 카메라가 그것을 따라간다. 배가 프레임 "
           "밖으로 나가는 자리에서 카메라가 들리면 언덕과 지게꾼이 들어온다"],
          ["지게 → 절문  (S10 → S11)", "인물 와이프",
           "지게를 진 일꾼이 카메라 앞을 가로질러 지나가며 화면을 한 번 덮는다. 그가 지나간 뒤 "
           "화면은 이미 절 문 앞이다"]],
         "세 방식이 물 → 시선 → 인물 로 갈리면서 마을의 시간이 자연스럽게 흐릅니다."),
        ("Q2", "“바닥 꽃들 같은 오브제가 움직임이 있는지”  (PPT S09)",
         [("현재는 전부 스틸입니다. ", True),
          ("27컷의 바닥 투사면이 모두 정지 이미지입니다. 움직임을 넣는 것을 권합니다 — 세 번째 "
           "면이 정지해 있으면 좌·우 화면만 움직일 때 바닥이 죽은 면으로 보입니다. 다만 전면 "
           "애니메이션은 과합니다. 바닥은 관객이 그 위를 걷는 면이라 시선을 끌면 안 됩니다.", False)],
         [["대상", "정도", "해당 컷"],
          ["풀 · 꽃", "바람에 아주 느리게 눕는 정도. 주기 8–12초의 완만한 파", "S4b · S5 · S6"],
          ["물", "잔물결과 반사의 흐름 — 이건 오히려 필수", "S7 · S8 · S9"],
          ["불", "도가니 빛의 일렁임", "S14 · S15 · S16"],
          ["빛줄기", "수면 아래 코스틱의 느린 이동", "S11 · S10 · S12"],
          ["돌 · 흙 · 마당", "정지가 자연스러운 바닥은 그대로 둡니다", "그 외"]],
         "지시 주시면 해당 바닥들을 루프 영상으로 다시 만들겠습니다."),
        ("Q3", "“S16에서 S17 어떻게 넘어갈까요?”  (저희 S8b → S9)",
         [("컷을 나누지 않고 하나로 잇습니다. ", True),
          ("큰 파도가 카메라를 향해 덮쳐오다가 물머리가 렌즈를 삼키는 순간, 화면 전체가 포말로 "
           "하얗게 찹니다. 그 흰 화면이 그대로 밝기를 잃으며 물속의 푸른 어둠으로 가라앉고, "
           "같은 순간 소리가 뚝 끊깁니다. 폭우와 파도의 굉음이 사라지고 수중의 먹먹함만 "
           "남습니다.", False)],
         None,
         "컷 전환이 아니라 카메라가 물에 잠기는 한 동작입니다. 관객도 같이 잠깁니다."),
        ("Q4", "“바닥에 무엇이 있을까요? 지금 좋은데 뭔가 있는 듯해서”  (PPT S17)",
         [("S17(저희 S9)의 바닥은 「수면 아래」입니다. ", True),
          ("물에 잠긴 직후의 시점이라 바닥면에는 위에서 내려다본 수면의 안쪽 — 은빛으로 "
           "일렁이는 수면 막, 빗방울이 위에서 때려 만드는 파문, 아래로 끌려 내려가는 기포와 "
           "부유물이 들어갑니다.\n"
           "즉 좌·우 화면이 “물에 잠긴 세계”를 보여줄 때, 바닥은 관객 머리 위에 있어야 할 "
           "수면이 발밑에 깔린 상태가 됩니다. 위아래가 뒤집힌 감각으로 침수를 체감시키려는 "
           "의도입니다.", False)],
         None,
         "“뭔가 있는 듯하다”고 느끼신 게 맞습니다. 다만 지금은 정지 이미지라 그 의도가 덜 "
         "읽힙니다. Q2 에서 말씀드린 움직임을 넣으면 가장 확실해지는 바닥이 이것입니다."),
    ]
    for tag, title, lead, tbl, tail in QA:
        sl = d.slide()
        y = d.head(sl, f"3 — 질문 4건   {tag}", title)
        d.text(sl, M, y, FULL, Emu(900000),
               [[(t, 11, b, ACCENT if b else INK) for t, b in lead]])
        y += Emu(400000) + Emu(int(160000 * (sum(len(t) for t, _ in lead) // 105 + 1)))
        if tbl:
            y = d.table(sl, M, y, FULL, tbl, [26, 24, 50], size=10)
            y += Emu(160000)
        img = {"Q3": "S08b_클라이막스_r21_va.png", "Q4": "S09_수면아래_건물없음_r20_va.png"}.get(tag)
        if img:
            y += d.pic(sl, canvas(img), M, y, FULL, None, px=2400) + Emu(220000)
        d.text(sl, M, y, FULL, Emu(500000), [[(tail, 11, True, ACCENT)]])

    # ══ 지시를 기다리는 것 ═══════════════════════════════════
    sl = d.slide()
    y = d.head(sl, "4", "지시를 기다리는 것")
    rows = [["", "사안", "선택지", "저희 의견"],
            ["1", "읍성 형식", "A. 조선 형식 그대로 / B. 고려 토축으로 변경",
             "A 권장 — 의도한 선택임을 기록"],
            ["2", "시장", "현행 유지 + 빠르게 통과 / 멍석·보따리 위주로 한 단계 낮추기",
             "현행 유지 + 빠른 통과"],
            ["3", "물속 유물", "보내주신 렌더 그대로 / 962년 기준 초기 청자·질그릇으로 교체",
             "지시 주시는 대로"],
            ["4", "바닥 움직임", "물·불·빛줄기 바닥을 루프 영상으로 만들지",
             "만드는 쪽 권장 (Q2)"]]
    y = d.table(sl, M, y, FULL, rows, [5, 18, 47, 30], size=11)
    d.text(sl, M, y + Emu(400000), FULL, Emu(700000), [
        [("그 외 진행 상황", 10, True, GREY)],
        [("· 27컷 전부 바닥 톤이 맞습니다 — 플레이트가 바뀐 컷은 바닥도 함께 다시 뽑았습니다", 10, False, INK)],
        [("· 여백 존(양 화면 사이 구간)은 27컷 전부 비어 있습니다", 10, False, INK)],
        [("· 확정 컷 원본은 무손실 PNG 로 보관 중이며, 컷 번호 순 파일명으로 전달 가능합니다", 10, False, INK)],
    ])

    # ══ 참고 출처 ═══════════════════════════════════════════
    sl = d.slide()
    y = d.head(sl, "", "참고 출처")
    srcs = [("『고려도경』", "ko.wikipedia.org/wiki/고려도경"),
            ("고려 복식 · 백저포", "우리역사넷 contents.history.go.kr"),
            ("청주 용두사지 철당간", "디지털청주문화대전 cheongju.grandculture.net"),
            ("준풍(峻豊) 연호", "디지털청주문화대전 cheongju.grandculture.net"),
            ("청주읍성 · 청주읍성도", "디지털청주문화대전 cheongju.grandculture.net"),
            ("읍성제 (토축 → 석축)", "문화콘텐츠닷컴 culturecontent.com"),
            ("장시의 성립", "우리역사넷 contents.history.go.kr"),
            ("시장의 역사 — 향시 · 허시 · 성읍시", "문화콘텐츠닷컴 culturecontent.com"),
            ("청자의 시작", "우리역사넷 contents.history.go.kr"),
            ("고려청자 편년", "ko.wikipedia.org/wiki/고려청자")]
    d.table(sl, M, y, FULL, [["자료", "출처"]] + [list(s) for s in srcs], [34, 66], size=11)
    d.foot(sl, "제오경 「염원 – 주성의 돛대」 · 2026 청주 국가유산 미디어아트 · 충청도병마절도사영문")

    dest = os.path.join(OUT, "제오경_2차피드백_회신.pptx")
    try:
        d.prs.save(dest)
    except PermissionError:
        stem, ext = os.path.splitext(dest)
        i = 2
        while os.path.exists(f"{stem}_{i}{ext}"):
            i += 1
        dest = f"{stem}_{i}{ext}"
        d.prs.save(dest)
        print("  ! 원래 파일이 파워포인트에 열려 있어 옆에 저장했다.", file=sys.stderr)
    print(f"회신 PPTX {len(d.prs.slides.__iter__.__self__._sldIdLst)}장  "
          f"{os.path.getsize(dest)/1024/1024:.1f} MB")
    print(f"  → {dest}")
    return 0


if __name__ == "__main__":
    sys.exit(build())
